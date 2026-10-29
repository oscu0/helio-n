import json
from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import h5py
import numpy as np
from openpyxl import load_workbook
import pandas as pd

from Library.SW.Report import write_event_stats_workbook
from Library.SW.Stats import (
    HSS_EVENT_REFERENCES,
    build_hss_yearly_stats,
    load_hss_event_catalogs,
    score_hss_events,
)

SCRIPT_SPEC = spec_from_file_location(
    "sw_stats_script", Path(__file__).resolve().parents[1] / "Scripts/Make/SW_Stats.py"
)
SW_STATS = module_from_spec(SCRIPT_SPEC)
SCRIPT_SPEC.loader.exec_module(SW_STATS)


def reference_event(start="2026-01-02", end="2026-01-02 03:00", peak=600.0):
    return pd.DataFrame({
        "T_start": [pd.Timestamp(start)], "T_end": [pd.Timestamp(end)], "v_max": [peak],
    })


class HSSEventStatsTests(unittest.TestCase):
    def setUp(self):
        self.events = reference_event()
        self.hours = pd.date_range("2026-01-01 12:00", "2026-01-02 15:00", freq="h")

    def test_threshold_and_no_ch_subset(self):
        for speed, hit, miss, no_ch in [
            (300.0, False, True, True),
            (300.00001, False, True, False),
            (400.0, False, True, False),
            (400.00001, True, False, False),
        ]:
            result = score_hss_events(self.events, pd.Series(speed, index=self.hours)).iloc[0]
            self.assertEqual((result.hit, result.miss, result.miss_no_ch), (hit, miss, no_ch))
            self.assertAlmostEqual(result.underreporting_km_s, 600.0 - speed)

    def test_padding_ties_and_one_hour_hit(self):
        prediction = pd.Series(300.0, index=self.hours)
        prediction.iloc[[0, -1]] = 500.0
        result = score_hss_events(self.events, prediction).iloc[0]
        self.assertTrue(result.hit)
        self.assertFalse(result.miss_no_ch)
        self.assertEqual(result.v_pred_peak, 500.0)
        self.assertEqual(result.t_pred_peak, self.hours[0])
        prediction.iloc[0] = 300.0
        self.assertTrue(score_hss_events(self.events, prediction).iloc[0].hit)

    def test_partial_context_uses_available_peak(self):
        prediction = pd.Series(500.0, index=self.hours)
        for incomplete in [prediction.iloc[1:], prediction.mask(prediction.index == self.hours[10])]:
            result = score_hss_events(self.events, incomplete).iloc[0]
            self.assertTrue(result.n_scored)
            self.assertTrue(result.hit)
            self.assertFalse(result.miss)
            self.assertEqual(result.v_pred_peak, 500.0)
            self.assertEqual(result.underreporting_km_s, 100.0)
        result = score_hss_events(self.events, pd.Series(300.0, index=self.hours[1:])).iloc[0]
        self.assertTrue(result.miss)
        self.assertFalse(result.miss_no_ch)

    def test_no_finite_prediction_is_unscored(self):
        for prediction in [pd.Series(dtype=float), pd.Series(np.nan, index=self.hours)]:
            result = score_hss_events(self.events, prediction).iloc[0]
            self.assertFalse(result.n_scored)
            self.assertFalse(result.hit)
            self.assertFalse(result.miss)
            self.assertFalse(result.miss_no_ch)
            self.assertTrue(pd.isna(result.v_pred_peak))
            self.assertTrue(pd.isna(result.underreporting_km_s))

    def test_empty_window_does_not_dilute_yearly_rates(self):
        events = pd.concat([
            self.events, reference_event("2026-01-04", "2026-01-04 03:00"),
        ], ignore_index=True)
        catalogs = {sheet: events for sheet, _sat, _filename in HSS_EVENT_REFERENCES}
        prediction = pd.Series(500.0, index=self.hours[1:])
        predictions = {sat: prediction for _sheet, sat, _filename in HSS_EVENT_REFERENCES}
        row = build_hss_yearly_stats(catalogs, predictions, "2026-01-01", "2027-01-01")["B1 - ACE"].iloc[0]
        self.assertEqual(row[["n_events", "n_scored", "n_unscored"]].tolist(), [2, 1, 1])
        self.assertEqual(row[["hit", "miss", "n_underreporting"]].tolist(), [1, 0, 1])
        self.assertEqual(row.hit_rate, 1.0)
        self.assertEqual(row.mean_underreporting_km_s, 100.0)

    def test_cross_year_aggregation_and_missing_reference_peak(self):
        events = pd.concat([
            reference_event("2025-12-31 23:00", "2026-01-01 01:00", 500.0),
            reference_event("2026-01-03", "2026-01-03 03:00", 700.0),
            reference_event("2026-01-05", "2026-01-05 03:00", 600.0),
            reference_event("2026-01-07", "2026-01-07 03:00", np.nan),
        ], ignore_index=True)
        prediction = pd.Series(600.0, index=pd.date_range("2025-12-30", "2026-01-08", freq="h"))
        prediction.loc["2026-01-02 12:00":"2026-01-03 15:00"] = 300.0
        prediction.loc["2026-01-04 12:00":"2026-01-05 15:00"] = 400.0
        catalogs = {sheet: events for sheet, _sat, _filename in HSS_EVENT_REFERENCES}
        predictions = {sat: prediction for _sheet, sat, _filename in HSS_EVENT_REFERENCES}
        tables = build_hss_yearly_stats(catalogs, predictions, "2025-12-31", "2026-01-08")
        rows = tables["B1 - ACE"].set_index("year")
        self.assertEqual(rows.loc[2025, "n_events"], 1)
        self.assertEqual(rows.loc[2025, "mean_underreporting_km_s"], -100.0)
        self.assertEqual(rows.loc[2026, ["hit", "miss", "miss_no_ch"]].tolist(), [1, 2, 1])
        self.assertEqual(rows.loc[2026, "n_underreporting"], 2)
        self.assertEqual(rows.loc[2026, "mean_underreporting_km_s"], 300.0)
        self.assertAlmostEqual(rows.loc[2026, "hit_rate"] + rows.loc[2026, "miss_rate"], 1.0)

    def test_catalog_selection_ignores_icme_flags(self):
        with TemporaryDirectory() as directory:
            original = pd.concat([
                reference_event("2025-12-31 23:00", "2026-01-01 01:00"), self.events,
            ], ignore_index=True)
            original["icme_flag"] = ["overlap", "near"]
            for _sheet, _sat, filename in HSS_EVENT_REFERENCES:
                original.to_csv(Path(directory) / filename, index=False)
            catalogs = load_hss_event_catalogs(directory, "2025-01-01", "2026-01-01")
            self.assertEqual(len(catalogs["B1 - ACE"]), 1)
            self.assertEqual(catalogs["B1 - ACE"].iloc[0].T_start.year, 2025)
            self.assertNotIn("icme_flag", catalogs["B1 - ACE"].columns)

    def test_missing_target_and_empty_reference(self):
        catalogs = {sheet: self.events.iloc[:0] for sheet, _sat, _filename in HSS_EVENT_REFERENCES}
        catalogs["B1 - ACE"] = self.events
        tables = build_hss_yearly_stats(catalogs, {}, "2026-01-01", "2027-01-01")
        self.assertEqual(tables["B1 - ACE"].iloc[0].n_unscored, 1)
        self.assertEqual(tables["B1 - PSP"].iloc[0].n_events, 0)
        self.assertTrue(pd.isna(tables["B1 - PSP"].iloc[0].hit_rate))
        with TemporaryDirectory() as directory:
            output = write_event_stats_workbook(tables, Path(directory) / "events.xlsx")
            book = load_workbook(output, data_only=True)
            self.assertEqual(book.sheetnames, [sheet for sheet, _sat, _filename in HSS_EVENT_REFERENCES])
            self.assertIsNone(book["B1 - PSP"]["H2"].value)
            self.assertEqual(book["B1 - ACE"].freeze_panes, "A2")
            self.assertEqual(book["B1 - ACE"]["H2"].number_format, "0.0%")

    def test_earth_samples_cube_not_native_ace(self):
        with TemporaryDirectory() as directory:
            cr = Path(directory) / "CR2306"
            cr.mkdir()
            hours = pd.date_range("2026-01-01", periods=3, freq="h")
            (cr / "manifest.json").write_text(json.dumps({
                "start_utc": str(hours[0]), "end_utc": str(hours[-1] + pd.Timedelta(hours=1)),
            }))
            speed = np.full((3, 2, 2), 999.0)
            speed[:, 0, 1] = [300.0, 450.0, 500.0]
            with h5py.File(cr / "cube.h5", "w") as cube:
                cube["time_ns"] = hours.asi8
                cube["phi"] = [0.0, 180.0]
                cube["r"] = [20.0, 215.0]
                cube["speed"] = speed
                cube["is_slow_wind"] = speed == 300.0
            pd.DataFrame({("satellite", "ace", "v_predict_raw"): [610.0, 620.0, 630.0]}, index=hours).to_parquet(cr / "series.parquet")
            predictions = SW_STATS.load_event_prediction_series(directory, hours[0], hours[-1] + pd.Timedelta(hours=1), ["ace"])
            self.assertEqual(predictions["earth"].tolist(), [300.0, 450.0, 500.0])
            self.assertEqual(predictions["ace"].tolist(), [610.0, 620.0, 630.0])

    def test_events_only_needs_no_original_observation_sources(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            cr = root / "archive/CR2306"
            cr.mkdir(parents=True)
            hours = pd.date_range("2025-12-31 12:00", "2026-01-01 15:00", freq="h")
            (cr / "manifest.json").write_text(json.dumps({
                "start_utc": str(hours[0]), "end_utc": str(hours[-1] + pd.Timedelta(hours=1)),
            }))
            speed = np.full((len(hours), 1, 1), 300.0)
            with h5py.File(cr / "cube.h5", "w") as cube:
                cube["time_ns"] = hours.asi8
                cube["phi"] = [0.0]
                cube["r"] = [215.0]
                cube["speed"] = speed
                cube["is_slow_wind"] = speed == 300.0
            pd.DataFrame({("satellite", "ace", "v_predict_raw"): 500.0}, index=hours).to_parquet(cr / "series.parquet")
            catalogs = root / "catalogs"
            catalogs.mkdir()
            for _sheet, _sat, filename in HSS_EVENT_REFERENCES:
                reference_event("2026-01-01", "2026-01-01 03:00").to_csv(catalogs / filename, index=False)
            output = root / "output"
            with patch.object(SW_STATS, "restore_observed_and_recurrent_series") as restore, patch.object(
                SW_STATS, "write_satellite_data_workbook"
            ) as write_hourly, patch.object(SW_STATS, "write_per_cr_stats_workbook") as write_cr, patch.object(
                SW_STATS, "write_summary_stats_workbook"
            ) as write_summary:
                result = SW_STATS.main([
                    "2306", "--events-only", "--archive-root", str(root / "archive"),
                    "--hss-list-dir", str(catalogs), "--output-dir", str(output),
                ])
            self.assertEqual(result, 0)
            restore.assert_not_called()
            write_hourly.assert_not_called()
            write_cr.assert_not_called()
            write_summary.assert_not_called()
            files = list(output.glob("*.xlsx"))
            self.assertEqual(len(files), 1)
            self.assertIn("(Event-based)", files[0].name)
            book = load_workbook(files[0], data_only=True)
            earth = list(book["Yermolaev checked - Earth"].values)
            ace = list(book["B1 - ACE"].values)
            self.assertEqual(earth[-1][5:7], (1, 1))
            self.assertEqual(ace[-1][4:7], (1, 0, 0))


if __name__ == "__main__":
    unittest.main()
